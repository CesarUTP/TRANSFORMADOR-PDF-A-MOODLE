"""
exportar_docx.py — el examen revisado como Word EDITABLE, con su clave (2.8).

Sale del MISMO plan que el PDF (`plan_impreso.py`): las mismas preguntas, el mismo orden y barajado, las mismas
letras impresas, la misma clave, las mismas indicaciones, fórmulas, fuentes y medidas. Aquí solo se traducen esos
bloques a Word: estilos de párrafo con nombre («Cátedra enunciado»…) que el docente puede cambiar, tablas reales,
numeración «Página X de Y» con campos de Word y un encabezado por sección.

    resultado = generar_docx(preguntas, answer_key, total_points, datos)
    resultado.docx, resultado.preguntas, resultado.avisos

Es un documento de APOYO como el PDF: lo que se importa a Moodle sigue siendo solo el XML.

Tipos de letra: Arial y Times New Roman se escriben con su nombre (el PDF usa Liberation Sans y Liberation Serif, que
miden lo mismo: las líneas se parten en los mismos sitios). DejaVu Sans no viene con Word: se INCRUSTA en el archivo
(Word de escritorio la respeta; Word web y Google Docs la sustituyen). Casillas y burbujas son imágenes pequeñas.

No hay forma de abrir el .docx en Word desde aquí: la prueba lee el archivo de vuelta (dev/test_exportar_docx.py).
"""

import html
import io
import re
import uuid
import zipfile
from dataclasses import dataclass, field, replace
from datetime import datetime, timezone
from typing import Any, Dict, List, Optional, Set, Tuple

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from lxml import etree
from PIL import Image as PILImage, ImageDraw, ImageFont
from reportlab.pdfbase import pdfmetrics

from plan_impreso import (
    FUENTES, GRIS, GRIS_CLARO, LADO_CASILLA, PASO_BURBUJAS, PASO_HUECOS, Burbujas, CuadroEstudiante, DatosExamen,
    Espacio, Estilo, FilaDocente, Grupo, Huecos, Imagen, LogosCabecera, Opcion, Parrafo, Plan, RejillaBurbujas,
    Renglones, Sangrado, Seccion, TablaHuecos, TablaPareo, _DIR_FUENTES, construir_plan,
)

# Cómo se llama cada tipo de letra en Word. DejaVu no viene con Word: se incrusta.
NOMBRE_WORD = {"dejavu": "DejaVu Sans", "arial": "Arial", "times": "Times New Roman"}
_RL_A_WORD = {"DejaVu": "DejaVu Sans", "LiberationSans": "Arial", "LiberationSerif": "Times New Roman"}
_INCRUSTADAS = {"DejaVu Sans": "dejavu"}
_ALINEAR = {0: WD_ALIGN_PARAGRAPH.LEFT, 1: WD_ALIGN_PARAGRAPH.CENTER, 2: WD_ALIGN_PARAGRAPH.RIGHT}
IDIOMA = "es-PA"


@dataclass
class ResultadoDocx:
    docx: bytes
    preguntas: int
    avisos: List[str] = field(default_factory=list)


# ── Leer el marcado del plan ────────────────────────────────────────────────

@dataclass
class Trozo:
    texto: Optional[str]              # None = salto de línea
    negrita: bool = False
    cursiva: bool = False
    fuente: Optional[str] = None      # nombre en ReportLab («DejaVu»…) si el trozo va con otra fuente
    tam: Optional[float] = None
    color: Optional[str] = None


_ETIQUETA = re.compile(r"<(/?)([a-zA-Z]+)([^>]*)>")
_ATRIBUTO = re.compile(r"""(\w+)\s*=\s*(?:'([^']*)'|"([^"]*)")""")
_BLANCOS = re.compile(r"[ \t\r\n\x0b\x0c]+")        # el espacio de no separación ( ) NO se colapsa, como en ReportLab


def leer_marcado(marcado: str) -> List[Trozo]:
    """El marcado de un párrafo del plan (`<b>`, `<i>`, `<br/>`, `<font name size color>`) como trozos de texto.
    Los blancos se juntan en uno y los de los extremos se quitan, igual que hace ReportLab al armar el PDF."""
    pila: List[Tuple[bool, bool, Optional[str], Optional[float], Optional[str]]] = [(False, False, None, None, None)]
    crudos: List[Trozo] = []
    pos = 0

    def texto(t: str):
        if t:
            n, c, f, tam, col = pila[-1]
            crudos.append(Trozo(html.unescape(t), n, c, f, tam, col))

    for m in _ETIQUETA.finditer(marcado):
        texto(marcado[pos:m.start()])
        pos = m.end()
        cierre, nombre, resto = m.group(1) == "/", m.group(2).lower(), m.group(3)
        if nombre == "br":
            crudos.append(Trozo(None))
        elif nombre in ("b", "i", "font"):
            if cierre:
                if len(pila) > 1:
                    pila.pop()
                continue
            n, c, f, tam, col = pila[-1]
            if nombre == "b":
                n = True
            elif nombre == "i":
                c = True
            else:
                for atr in _ATRIBUTO.finditer(resto):
                    clave, valor = atr.group(1).lower(), atr.group(2) if atr.group(2) is not None else atr.group(3)
                    if clave == "name":
                        f = valor
                    elif clave == "size":
                        tam = float(valor)
                    elif clave == "color":
                        col = valor
            pila.append((n, c, f, tam, col))
    texto(marcado[pos:])

    salida: List[Trozo] = []
    ultimo_blanco = True
    for t in crudos:
        if t.texto is None:
            salida.append(t)
            ultimo_blanco = True
            continue
        txt = _BLANCOS.sub(" ", t.texto)
        if ultimo_blanco and txt.startswith(" "):
            txt = txt[1:]
        if not txt:
            continue
        ultimo_blanco = txt.endswith(" ")
        salida.append(replace(t, texto=txt))
    for k in range(len(salida) - 1, -1, -1):
        if salida[k].texto is None:
            break
        recortado = salida[k].texto.rstrip(" ")
        salida[k] = replace(salida[k], texto=recortado)
        if recortado:
            break
    return [t for t in salida if t.texto != ""]


# ── Imágenes de casillas y burbujas ─────────────────────────────────────────

_PX = 12                                 # píxeles por punto: se ven nítidas al ampliar o imprimir
_CACHE_IMG: Dict[tuple, bytes] = {}


def _png(im: PILImage.Image) -> bytes:
    b = io.BytesIO()
    im.save(b, "PNG")
    return b.getvalue()


def imagen_casilla(redonda: bool) -> bytes:
    """El círculo (una respuesta) o cuadrado (varias) de una opción."""
    clave = ("casilla", redonda)
    if clave not in _CACHE_IMG:
        lado = round(LADO_CASILLA * _PX)
        im = PILImage.new("RGBA", (lado, lado), (255, 255, 255, 0))
        d = ImageDraw.Draw(im)
        g = round(0.8 * _PX)
        caja = (g // 2, g // 2, lado - 1 - g // 2, lado - 1 - g // 2)
        (d.ellipse if redonda else d.rectangle)(caja, outline=(0, 0, 0, 255), width=g)
        _CACHE_IMG[clave] = _png(im)
    return _CACHE_IMG[clave]


def imagen_burbuja(letra: str, redonda: bool, lleno: bool, archivo_fuente: str) -> bytes:
    """Una burbuja de la hoja de respuestas con su letra (rellena en la clave)."""
    clave = ("burbuja", letra, redonda, lleno, archivo_fuente)
    if clave not in _CACHE_IMG:
        lado = round(13.2 * _PX)
        im = PILImage.new("RGBA", (lado, lado), (255, 255, 255, 0))
        d = ImageDraw.Draw(im)
        g = round(0.8 * _PX)
        caja = (g // 2, g // 2, lado - 1 - g // 2, lado - 1 - g // 2)
        relleno = (0, 0, 0, 255) if lleno else (255, 255, 255, 255)
        (d.ellipse if redonda else d.rectangle)(caja, outline=(0, 0, 0, 255), width=g, fill=relleno)
        fuente = ImageFont.truetype(str(_DIR_FUENTES / archivo_fuente), round(7 * _PX))
        d.text((lado / 2, lado / 2), letra, font=fuente, fill=(255, 255, 255, 255) if lleno else (0, 0, 0, 255), anchor="mm")
        _CACHE_IMG[clave] = _png(im)
    return _CACHE_IMG[clave]


# ── Utilidades de XML de Word ───────────────────────────────────────────────

_ORDEN_TBLPR = ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize", "tblStyleColBandSize", "tblW", "jc",
                "tblCellSpacing", "tblInd", "tblBorders", "shd", "tblLayout", "tblCellMar", "tblLook"]
_ORDEN_TCPR = ["cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd", "noWrap", "tcMar", "textDirection",
               "tcFitText", "vAlign", "hideMark"]
_ORDEN_PPR = ["pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl", "numPr", "suppressLineNumbers",
              "pBdr", "shd", "tabs", "suppressAutoHyphens", "kinsoku", "wordWrap", "overflowPunct", "topLinePunct", "autoSpaceDE",
              "autoSpaceDN", "bidi", "adjustRightInd", "snapToGrid", "spacing", "ind", "contextualSpacing", "mirrorIndents",
              "suppressOverlap", "jc", "textDirection", "textAlignment", "textboxTightWrap", "outlineLvl", "divId", "cnfStyle",
              "rPr", "sectPr", "pPrChange"]


def _ordenar(el, orden: List[str]) -> None:
    """Deja los hijos de `el` en el orden que pide el esquema de Word (si no, Word puede rechazar el archivo)."""
    def clave(h):
        nombre = etree.QName(h).localname
        return orden.index(nombre) if nombre in orden else len(orden)
    hijos = sorted(list(el), key=clave)
    for h in hijos:
        el.remove(h)
    for h in hijos:
        el.append(h)


def _twips(pt: float) -> str:
    return str(round(pt * 20))


def _hijo(el, nombre: str, **atributos):
    previo = el.find(qn(nombre))
    if previo is not None:
        el.remove(previo)
    nuevo = OxmlElement(nombre)
    for k, v in atributos.items():
        nuevo.set(qn("w:" + k), str(v))
    el.append(nuevo)
    return nuevo


def _fuente_rpr(rpr, nombre: str) -> None:
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.insert(0, rf)
    for k in ("ascii", "hAnsi", "cs", "eastAsia"):
        rf.set(qn("w:" + k), nombre)
    for k in ("asciiTheme", "hAnsiTheme", "cstheme", "eastAsiaTheme"):
        if rf.get(qn("w:" + k)) is not None:
            del rf.attrib[qn("w:" + k)]


# ── El documento ────────────────────────────────────────────────────────────

class _Docx:
    def __init__(self, plan: Plan):
        self.plan, self.geo = plan, plan.geo
        self.doc = Document()
        self.word_e = NOMBRE_WORD[plan.clave_fuente_e]
        self.word_p = NOMBRE_WORD[plan.clave_fuente_p]
        self.usos: Set[Tuple[str, bool, bool]] = set()          # (fuente, negrita, cursiva) de las incrustables que se usan
        self._estilos: Dict[Estilo, str] = {}
        self._nombres: Set[str] = set()
        self.cont = 0.0                                          # sangría que dan los bloques `Sangrado` en curso
        self.prev_despues = 0.0                                  # «espacio después» del párrafo anterior (se funde con el «antes»)
        self._ultimo: Optional[Tuple[str, Any, float]] = None    # ("p", párrafo, su espacio después) o ("t", tabla, 0)
        self._grupo: Optional[list] = None                       # elementos del `Grupo` en curso
        self._preparar_documento()

    # ---- el documento en blanco ----
    def _preparar_documento(self) -> None:
        d = self.doc
        normal = d.styles["Normal"]
        self._fuente_estilo(normal, self.word_p, False, False)
        normal.font.size = Pt(self.plan.tam_preguntas)
        pf = normal.paragraph_format
        pf.space_before, pf.space_after = Pt(0), Pt(0)
        pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
        pf.widow_control = False
        # Valores de partida de la plantilla de python-docx: idioma español y nada de espacios extra entre párrafos.
        defaults = d.styles.element.find(qn("w:docDefaults"))
        rpr = defaults.find(qn("w:rPrDefault")).find(qn("w:rPr"))
        lang = rpr.find(qn("w:lang"))
        if lang is not None:
            lang.set(qn("w:val"), IDIOMA)
            lang.set(qn("w:eastAsia"), IDIOMA)
        sp = defaults.find(qn("w:pPrDefault")).find(qn("w:pPr")).find(qn("w:spacing"))
        if sp is not None:
            sp.set(qn("w:after"), "0")
            sp.set(qn("w:line"), "240")
        # Modo de compatibilidad de Word actual (sin esto Word muestra «Modo de compatibilidad»).
        for cs in d.settings.element.iter(qn("w:compatSetting")):
            if cs.get(qn("w:name")) == "compatibilityMode":
                cs.set(qn("w:val"), "15")
        cp = d.core_properties
        cp.title, cp.author, cp.subject = self.plan.meta["titulo"], self.plan.meta["autor"], self.plan.meta["materia"]
        cp.last_modified_by, cp.comments, cp.keywords = "Cátedra", "", ""
        cp.created = cp.modified = datetime.now(timezone.utc).replace(tzinfo=None)

    def _fuente_estilo(self, estilo, nombre_word: str, negrita: bool, cursiva: bool) -> None:
        estilo.font.name = nombre_word
        _fuente_rpr(estilo.element.get_or_add_rPr(), nombre_word)
        self._usar(nombre_word, negrita, cursiva)

    def _usar(self, nombre_word: str, negrita: bool, cursiva: bool) -> None:
        if nombre_word in _INCRUSTADAS:
            self.usos.add((nombre_word, negrita, cursiva))

    # ---- estilos de párrafo ----
    def _estilo(self, e: Estilo) -> str:
        if e in self._estilos:
            return self._estilos[e]
        base = f"Cátedra {e.nombre or 'texto'}"
        nombre, n = base, 1
        while nombre in self._nombres:
            n += 1
            nombre = f"{base} {n}"
        self._nombres.add(nombre)
        st = self.doc.styles.add_style(nombre, WD_STYLE_TYPE.PARAGRAPH)
        st.base_style = self.doc.styles["Normal"]
        st.quick_style = True
        self._fuente_estilo(st, self.word_e if e.grupo == "e" else self.word_p, e.negrita, e.cursiva)
        st.font.size = Pt(e.tam)
        if e.negrita:
            st.font.bold = True
        if e.cursiva:
            st.font.italic = True
        if e.color:
            st.font.color.rgb = RGBColor.from_string(e.color.lstrip("#").upper())
        pf = st.paragraph_format
        pf.alignment = _ALINEAR[e.alinear]
        pf.left_indent = Pt(e.sangria)
        pf.space_before, pf.space_after = Pt(e.antes), Pt(e.despues)
        pf.line_spacing = Pt(e.leading)
        pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        pf.widow_control = False
        self._estilos[e] = nombre
        return nombre

    def _fuente_de(self, e: Estilo) -> str:
        return self.word_e if e.grupo == "e" else self.word_p

    # ---- contenedores (el cuerpo o una celda) ----
    @staticmethod
    def _sin_usar(celda) -> bool:
        """La celda sigue con el párrafo vacío que Word le pone al crearla."""
        ps = celda.paragraphs
        return len(ps) == 1 and not ps[0].runs and not ps[0].text

    def _nuevo(self, destino) -> Any:
        """Un párrafo nuevo en el cuerpo o en una celda (la primera vez, el que la celda ya trae vacío)."""
        if destino is self.doc:
            return self.doc.add_paragraph()
        if self._sin_usar(destino):
            return destino.paragraphs[0]
        return destino.add_paragraph()

    # ---- texto ----
    def _run(self, par, texto: str, e: Estilo, t: Optional[Trozo] = None, negrita: Optional[bool] = None,
             tam: Optional[float] = None):
        r = par.add_run(texto)
        neg = e.negrita or bool(t and t.negrita) or bool(negrita)
        cur = e.cursiva or bool(t and t.cursiva)
        if (t and t.negrita and not e.negrita) or negrita:
            r.bold = True
        if t and t.cursiva and not e.cursiva:
            r.italic = True
        fuente = self._fuente_de(e)
        if t and t.fuente:
            fuente = _RL_A_WORD.get(t.fuente, fuente)
            _fuente_rpr(r._r.get_or_add_rPr(), fuente)
        if (t and t.tam) or tam:
            r.font.size = Pt(tam or t.tam)
        if t and t.color:
            r.font.color.rgb = RGBColor.from_string(t.color.lstrip("#").upper())
        self._usar(fuente, neg, cur)
        return r

    def _contenido(self, par, marcado: str, e: Estilo) -> None:
        for t in leer_marcado(marcado):
            if t.texto is None:
                par.add_run().add_break()
            else:
                self._run(par, t.texto, e, t)

    def _formato(self, par, e: Estilo, vineta: bool, antes: Optional[float] = None, cont: Optional[float] = None) -> None:
        """`cont`: la sangría del bloque en curso (en una celda no se aplica: la tabla ya lleva la suya)."""
        c = self.cont if cont is None else cont
        pf = par.paragraph_format
        if vineta:
            pf.left_indent = Pt(e.sangria + c)
            pf.first_line_indent = Pt(-(e.sangria - e.vineta_sangria))
        elif c:
            pf.left_indent = Pt(e.sangria + c)
        if antes is not None and abs(antes - e.antes) > 1e-9:
            pf.space_before = Pt(antes)
        pf.keep_together = True

    def _parrafo(self, p: Parrafo, destino=None) -> Any:
        """Un párrafo del plan. En el cuerpo, el «espacio antes» se funde con el «después» del anterior (como en el PDF)."""
        destino = destino or self.doc
        e = p.estilo
        par = self._nuevo(destino)
        par.style = self.doc.styles[self._estilo(e)]
        en_cuerpo = destino is self.doc
        antes = max(e.antes - self.prev_despues, 0.0) if en_cuerpo else None
        self._formato(par, e, p.vineta is not None, antes, self.cont if en_cuerpo else 0.0)
        if p.vineta is not None:
            self._run(par, p.vineta, e, negrita=e.vineta_negrita, tam=e.vineta_tam)
            par.add_run("\t")
        self._contenido(par, p.marcado, e)
        if en_cuerpo:
            self.prev_despues = e.despues
            self._registrar("p", par, e.despues)
        return par

    def _registrar(self, tipo: str, objeto: Any, despues: float = 0.0) -> None:
        self._ultimo = (tipo, objeto, despues)
        if self._grupo is not None:
            self._grupo.append((tipo, objeto))

    @staticmethod
    def _foto(run, datos: bytes, ancho: float, alto: float) -> None:
        """Una imagen en línea. Sin distancias explícitas, LibreOffice le pone 9 pt de aire a cada lado y Word 0."""
        run.add_picture(io.BytesIO(datos), width=Pt(ancho), height=Pt(alto))
        for inline in run._r.iter("{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}inline"):
            for a in ("distT", "distB", "distL", "distR"):
                inline.set(a, "0")

    # ---- bloques ----
    def bloques(self, bloques: list) -> None:
        for b in bloques:
            self.bloque(b)

    def bloque(self, b: Any) -> None:
        if isinstance(b, Parrafo):
            self._parrafo(b)
        elif isinstance(b, Espacio):
            self._espacio(b.alto)
        elif isinstance(b, Imagen):
            self._imagen_suelta(b)
        elif isinstance(b, Sangrado):
            self.cont += b.izquierda
            self.bloques(b.hijos)
            self.cont -= b.izquierda
        elif isinstance(b, Grupo):
            self._grupo_en_curso(b)
        elif isinstance(b, Opcion):
            self._opcion(b)
        elif isinstance(b, TablaPareo):
            self._tabla_pareo(b)
        elif isinstance(b, TablaHuecos):
            self._tabla_huecos(b)
        elif isinstance(b, Renglones):
            self._renglones(b)
        elif isinstance(b, RejillaBurbujas):
            self._rejilla(b)
        elif isinstance(b, Huecos):
            self._huecos(b)
        elif isinstance(b, LogosCabecera):
            self._logos(b)
        elif isinstance(b, CuadroEstudiante):
            self._cuadro(b)
        elif isinstance(b, FilaDocente):
            self._fila_docente(b)
        else:
            raise TypeError(f"bloque desconocido: {type(b).__name__}")

    def _grupo_en_curso(self, g: Grupo) -> None:
        """Lo de un grupo no se separa entre páginas mientras quepa: cada párrafo (y fila) se une con el siguiente."""
        previo, self._grupo = self._grupo, []
        self.bloques(g.hijos)
        elementos, self._grupo = self._grupo, previo
        for i, (tipo, obj) in enumerate(elementos):
            ultimo = i == len(elementos) - 1
            if tipo == "p":
                obj.paragraph_format.keep_with_next = not ultimo
            else:
                filas = obj.rows
                for j, fila in enumerate(filas):
                    for celda in fila.cells:
                        for par in celda.paragraphs:
                            par.paragraph_format.keep_with_next = not (ultimo and j == len(filas) - 1)
                            par.paragraph_format.keep_together = True
        if previo is not None:
            previo.extend(elementos)

    def _espacio(self, alto: float) -> None:
        """Un espacio vertical: se suma al «después» del párrafo anterior o, si no hay (o es una tabla), es un párrafo vacío."""
        if self._ultimo and self._ultimo[0] == "p":
            _tipo, par, despues = self._ultimo
            par.paragraph_format.space_after = Pt(despues + alto)
            self._ultimo = ("p", par, despues + alto)
        else:
            par = self.doc.add_paragraph()
            self._formato_vacio(par, alto)
            self._registrar("p", par, 0.0)
        self.prev_despues = 0.0

    def _formato_vacio(self, par, alto: float) -> None:
        pf = par.paragraph_format
        pf.space_before, pf.space_after = Pt(0), Pt(0)
        pf.line_spacing = Pt(alto)
        pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        r = par.add_run()
        r.font.size = Pt(1)

    def _imagen_suelta(self, im: Imagen, destino=None) -> Any:
        destino = destino or self.doc
        par = self._nuevo(destino)
        pf = par.paragraph_format
        pf.space_before, pf.space_after = Pt(0), Pt(0)
        pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
        pf.alignment = WD_ALIGN_PARAGRAPH.RIGHT if im.alinear == "RIGHT" else WD_ALIGN_PARAGRAPH.LEFT
        if self.cont and destino is self.doc:
            pf.left_indent = Pt(self.cont)
        pf.keep_together = True
        self._foto(par.add_run(), im.datos, im.ancho, im.alto)
        if destino is self.doc:
            self.prev_despues = 0.0
            self._registrar("p", par, 0.0)
        return par

    def _opcion(self, o: Opcion) -> None:
        """Una opción: su casilla (imagen) y el texto con sangría francesa, como la tabla del PDF."""
        e = self.plan.estilos["opcion"]
        par = self.doc.add_paragraph(style=self._estilo(e))
        self._formato(par, e, False)
        pf = par.paragraph_format
        pf.left_indent = Pt(e.sangria + self.cont + 18)
        pf.first_line_indent = Pt(-18)
        pf.space_before = Pt(0)
        pf.space_after = Pt(3)
        im = par.add_run()
        self._foto(im, imagen_casilla(o.casilla == "redonda"), LADO_CASILLA, LADO_CASILLA)
        im._r.get_or_add_rPr().append(self._posicion(-2))
        par.add_run("\t")
        self._contenido(par, o.marcado, e)
        self.prev_despues = 0.0          # en el PDF es el relleno de la fila: lo que sigue conserva todo su «espacio antes»
        self._registrar("p", par, 3.0)

    @staticmethod
    def _posicion(medios_puntos: int):
        el = OxmlElement("w:position")
        el.set(qn("w:val"), str(medios_puntos))
        return el

    def _renglones(self, r: Renglones) -> None:
        """Los renglones de una respuesta de desarrollo: una tabla de filas con su raya (así cada raya sale siempre).
        Como en el PDF, la 1.ª raya queda a 4 pt del borde de arriba y sobran 4 pt menos que un paso debajo de la última."""
        ancho = self.geo.ancho - self.geo.sangria
        alturas = [4.0] + [r.paso] * (r.n - 1) + [r.paso - 4.0]
        t = self._tabla([ancho], len(alturas), indent=self.cont)
        for i, (fila, alto) in enumerate(zip(t.rows, alturas)):
            fila.height, fila.height_rule = Pt(alto), WD_ROW_HEIGHT_RULE.EXACTLY
            if i < r.n:
                self._borde(fila.cells[0], bottom=(0.6, GRIS_CLARO))
            self._vacio(fila.cells[0])
        self._registrar("t", t)
        self.prev_despues = 0.0

    def _vacio(self, celda) -> None:
        par = celda.paragraphs[0]
        self._formato_vacio(par, 1)

    # ---- tablas ----
    def _tabla(self, anchos: List[float], filas: int = 1, indent: float = 0.0, margen: Tuple[float, float, float, float] = (0, 0, 0, 0)):
        """Una tabla de ancho fijo y sin bordes. `margen` = (arriba, derecha, abajo, izquierda) de cada celda, en pt."""
        if self._ultimo and self._ultimo[0] == "t":
            sep = self.doc.add_paragraph()                  # dos tablas pegadas se funden en una: se separan
            self._formato_vacio(sep, 1)
        t = self.doc.add_table(rows=filas, cols=len(anchos))
        t.autofit = False
        pr = t._tbl.tblPr
        _hijo(pr, "w:tblW", w=_twips(sum(anchos)), type="dxa")
        _hijo(pr, "w:tblInd", w=_twips(indent), type="dxa")
        mar = _hijo(pr, "w:tblCellMar")
        for lado, v in zip(("top", "left", "bottom", "right"), (margen[0], margen[3], margen[2], margen[1])):
            e = OxmlElement("w:" + lado)
            e.set(qn("w:w"), _twips(v))
            e.set(qn("w:type"), "dxa")
            mar.append(e)
        _ordenar(pr, _ORDEN_TBLPR)
        for i, w in enumerate(anchos):
            t.columns[i].width = Pt(w)
            for fila in t.rows:
                fila.cells[i].width = Pt(w)
        for fila in t.rows:
            trpr = fila._tr.get_or_add_trPr()
            if trpr.find(qn("w:cantSplit")) is None:
                trpr.insert(0, OxmlElement("w:cantSplit"))
        self._ultimo = ("t", t, 0.0)
        return t

    @staticmethod
    def _borde(celda, **lados) -> None:
        """Bordes de una celda: lado=(grosor en pt, color «#rrggbb») con lado en top/left/bottom/right."""
        pr = celda._tc.get_or_add_tcPr()
        b = pr.find(qn("w:tcBorders"))
        if b is None:
            b = OxmlElement("w:tcBorders")
            pr.append(b)
        for lado in ("top", "left", "bottom", "right"):
            if lado in lados:
                grosor, color = lados[lado]
                e = OxmlElement("w:" + lado)
                e.set(qn("w:val"), "single")
                e.set(qn("w:sz"), str(max(2, round(grosor * 8))))
                e.set(qn("w:space"), "0")
                e.set(qn("w:color"), color.lstrip("#").upper())
                b.append(e)
        _ordenar(b, ["top", "left", "start", "bottom", "right", "end", "insideH", "insideV"])
        _ordenar(pr, _ORDEN_TCPR)

    @staticmethod
    def _vertical(celda, donde: str) -> None:
        pr = celda._tc.get_or_add_tcPr()
        v = pr.find(qn("w:vAlign"))
        if v is None:
            v = OxmlElement("w:vAlign")
            pr.append(v)
        v.set(qn("w:val"), donde)
        _ordenar(pr, _ORDEN_TCPR)

    def _tabla_pareo(self, b: TablaPareo) -> None:
        mitad = (self.geo.ancho - self.geo.sangria) / 2
        t = self._tabla([mitad, mitad], len(b.filas), indent=self.cont, margen=(3, 6, 4, 0))
        for i, (izq, der) in enumerate(b.filas):
            for j, p in enumerate((izq, der)):
                self._parrafo(p, t.rows[i].cells[j])
                self._vertical(t.rows[i].cells[j], "top")
                if i == 0:
                    self._borde(t.rows[i].cells[j], bottom=(0.5, GRIS_CLARO))
        self._registrar("t", t)
        self.prev_despues = 0.0

    def _tabla_huecos(self, b: TablaHuecos) -> None:
        g = self.geo
        t = self._tabla([28.35, g.ancho - g.sangria - 28.35], len(b.filas), indent=self.cont, margen=(5, 0, 5, 0))
        for i, (num, celda) in enumerate(b.filas):
            self._parrafo(num, t.rows[i].cells[0])
            for p in celda:
                self._parrafo(p, t.rows[i].cells[1])
            for c in t.rows[i].cells:
                self._vertical(c, "top")
                if i < len(b.filas) - 1:
                    self._borde(c, bottom=(0.4, GRIS_CLARO))
        self._registrar("t", t)
        self.prev_despues = 0.0

    def _rejilla(self, b: RejillaBurbujas) -> None:
        g = self.geo
        t = self._tabla([g.ancho / b.columnas] * b.columnas, len(b.celdas), margen=(1, 8, 1, 0))
        fe = FUENTES[self.plan.clave_fuente_p][1][0]
        e = Estilo("p", 9, 19, negrita=True, nombre="burbujas")
        for i, fila in enumerate(b.celdas):
            for j, bur in enumerate(fila):
                celda = t.rows[i].cells[j]
                if bur is None:
                    self._vacio(celda)
                    continue
                par = self._nuevo(celda)
                par.style = self.doc.styles[self._estilo(e)]
                par.paragraph_format.keep_together = True
                etiqueta = f"{bur.num}."
                x0 = max(24.0, pdfmetrics.stringWidth(etiqueta, self.plan.fuente_p + "-B", 9) + 6)
                for k in range(len(bur.letras)):
                    par.paragraph_format.tab_stops.add_tab_stop(Pt(x0 + k * PASO_BURBUJAS), WD_TAB_ALIGNMENT.LEFT)
                self._run(par, etiqueta, e)
                for letra in bur.letras:
                    par.add_run("\t")
                    r = par.add_run()
                    self._foto(r, imagen_burbuja(letra, bur.redonda, letra in bur.marcadas, fe), 13.2, 13.2)
                    r._r.get_or_add_rPr().append(self._posicion(-4))
        self._registrar("t", t)
        self.prev_despues = 0.0

    def _huecos(self, h: Huecos) -> None:
        """«6.  1. ______  2. ______»: cada línea es un párrafo con tabuladores; la raya es el relleno del tabulador."""
        e = Estilo("p", h.tam, PASO_HUECOS, negrita=True, nombre="respuestas")
        for r, fila in enumerate(h.lineas):
            par = self.doc.add_paragraph(style=self._estilo(e))
            pf = par.paragraph_format
            pf.keep_together = True
            if self.cont:
                pf.left_indent = Pt(self.cont)
            if r == 0:
                self._run(par, f"{h.num}.", e)
            for x, etiqueta, valor, w_et, largo, tam_v in fila:
                fin = x + w_et + largo
                pf.tab_stops.add_tab_stop(Pt(x), WD_TAB_ALIGNMENT.LEFT)
                par.add_run("\t")
                if etiqueta:
                    self._run(par, etiqueta, e)
                    pf.tab_stops.add_tab_stop(Pt(x + w_et), WD_TAB_ALIGNMENT.LEFT)
                    par.add_run("\t")
                if valor:
                    pf.tab_stops.add_tab_stop(Pt(x + w_et + largo / 2), WD_TAB_ALIGNMENT.CENTER, WD_TAB_LEADER.LINES)
                    par.add_run("\t")
                    r_valor = self._run(par, valor, e, tam=tam_v)
                    r_valor.underline = True
                pf.tab_stops.add_tab_stop(Pt(fin), WD_TAB_ALIGNMENT.LEFT, WD_TAB_LEADER.LINES)
                par.add_run("\t")
            self.prev_despues = 0.0
            self._registrar("p", par, 0.0)

    # ---- encabezado de la primera hoja ----
    def _logos(self, b: LogosCabecera) -> None:
        g = self.geo
        t = self._tabla([b.lado, g.ancho - 2 * b.lado, b.lado])
        izq, centro, der = t.rows[0].cells
        for celda in (izq, centro, der):
            self._vertical(celda, "center")
        if b.izquierdo:
            self._imagen_suelta(b.izquierdo, izq)
        for p in b.centro:
            self._parrafo(p, centro)
        if b.derecho:
            self._imagen_suelta(b.derecho, der)
        for celda in (izq, der):
            if self._sin_usar(celda):
                self._vacio(celda)
        self._registrar("t", t)
        self.prev_despues = 0.0

    def _cuadro(self, b: CuadroEstudiante) -> None:
        t = self._tabla(b.anchos, 2, margen=(0, 0, 0, 6))
        for i, fila in enumerate(b.filas):
            t.rows[i].height, t.rows[i].height_rule = Pt(b.alto_fila), WD_ROW_HEIGHT_RULE.EXACTLY
            for j, c in enumerate(fila):
                celda = t.rows[i].cells[j]
                self._vertical(celda, "center")
                self._borde(celda, top=(0.8, "#000000"), left=(0.8, "#000000"), bottom=(0.8, "#000000"), right=(0.8, "#000000"))
                if isinstance(c, Parrafo):
                    self._parrafo(c, celda)
                else:
                    self._vacio(celda)
        self._registrar("t", t)
        self.prev_despues = 0.0

    def _fila_docente(self, b: FilaDocente) -> None:
        anchos = [b.ancho_izquierda] + ([b.ancho_nota] if b.nota else [])
        t = self._tabla(anchos)
        for j, p in enumerate([b.izquierda, b.nota][:len(anchos)]):
            celda = t.rows[0].cells[j]
            self._vertical(celda, "bottom")
            if p is not None:
                self._parrafo(p, celda)
            else:
                self._vacio(celda)
        self._registrar("t", t)
        self.prev_despues = 0.0

    # ---- secciones, encabezados y pies ----
    def _campo(self, par, instruccion: str, e: Estilo, valor: str = "1") -> None:
        def run(hijo):
            r = par.add_run()
            r._r.append(hijo)
            return r
        c = OxmlElement("w:fldChar")
        c.set(qn("w:fldCharType"), "begin")
        run(c)
        t = OxmlElement("w:instrText")
        t.set(qn("xml:space"), "preserve")
        t.text = f" {instruccion} "
        run(t)
        c = OxmlElement("w:fldChar")
        c.set(qn("w:fldCharType"), "separate")
        run(c)
        par.add_run(valor)
        c = OxmlElement("w:fldChar")
        c.set(qn("w:fldCharType"), "end")
        run(c)

    def _pie(self, parte) -> None:
        e = Estilo("e", 8, 10, color=GRIS, alinear=2, nombre="pie")
        parte.is_linked_to_previous = False
        par = parte.paragraphs[0]
        par.style = self.doc.styles[self._estilo(e)]
        self._run(par, "Página ", e)
        self._campo(par, "PAGE", e)
        self._run(par, " de ", e)
        self._campo(par, "SECTIONPAGES", e)

    def _encabezado(self, parte, s: Seccion) -> None:
        e = Estilo("e", 8, 10, color=GRIS, nombre="encabezado corrido")
        parte.is_linked_to_previous = False
        par = parte.paragraphs[0]
        par.style = self.doc.styles[self._estilo(e)]
        pf = par.paragraph_format
        pf.tab_stops.add_tab_stop(Pt(self.geo.ancho), WD_TAB_ALIGNMENT.RIGHT)
        self._run(par, s.izquierda, e)
        par.add_run("\t")
        self._run(par, s.derecha, e)
        ppr = par._p.get_or_add_pPr()
        bd = OxmlElement("w:pBdr")
        linea = OxmlElement("w:bottom")
        for k, v in (("val", "single"), ("sz", "4"), ("space", "4"), ("color", GRIS_CLARO.lstrip("#").upper())):
            linea.set(qn("w:" + k), v)
        bd.append(linea)
        ppr.append(bd)
        _ordenar(ppr, _ORDEN_PPR)

    def _vacia_primera(self, parte) -> None:
        parte.is_linked_to_previous = False
        par = parte.paragraphs[0]
        self._formato_vacio(par, 1)

    def _seccion(self, s: Seccion, primera: bool):
        g = self.geo
        sec = self.doc.sections[0] if primera else self.doc.add_section(WD_SECTION.NEW_PAGE)
        sec.page_width, sec.page_height = Pt(g.papel[0]), Pt(g.papel[1])
        sec.left_margin, sec.right_margin = Pt(g.m_izq), Pt(g.m_der)
        sec.top_margin, sec.bottom_margin = Pt(g.m_sup), Pt(g.m_inf)
        y_enc = max(0.55 * 28.3465, g.m_sup - 0.8 * 28.3465)
        sec.header_distance = Pt(max(10.0, y_enc - 7.5))
        sec.footer_distance = Pt(max(8.0, max(0.45 * 28.3465, g.m_inf - 0.8 * 28.3465) - 2))
        sec.different_first_page_header_footer = s.primera_sin_encabezado
        # Cada sección numera sus páginas desde 1 («Página X de Y» propio).
        pr = sec._sectPr
        for viejo in pr.findall(qn("w:pgNumType")):
            pr.remove(viejo)
        num = OxmlElement("w:pgNumType")
        num.set(qn("w:start"), "1")
        pr.find(qn("w:pgMar")).addnext(num)
        self._encabezado(sec.header, s)
        self._pie(sec.footer)
        if s.primera_sin_encabezado:
            self._vacia_primera(sec.first_page_header)
            self._pie(sec.first_page_footer)
        return sec

    def construir(self) -> bytes:
        for i, s in enumerate(self.plan.secciones):
            self._seccion(s, i == 0)
            self._ultimo = None
            self.prev_despues = 0.0
            self.bloques(s.bloques)
        buf = io.BytesIO()
        self.doc.save(buf)
        return _rematar(buf.getvalue(), self.usos)


# ── Retoques al archivo ya guardado (fuentes incrustadas, miniatura) ────────

_NS_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_NS_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_REL_FUENTE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/font"
_REL_MINIATURA = "http://schemas.openxmlformats.org/package/2006/relationships/metadata/thumbnail"
_TIPO_FUENTE = "application/vnd.openxmlformats-officedocument.obfuscatedFont"
_ESTILOS_ARCHIVO = {(False, False): 0, (True, False): 1, (False, True): 2, (True, True): 3}
_ETIQ_INCRUSTAR = {(False, False): "embedRegular", (True, False): "embedBold", (False, True): "embedItalic", (True, True): "embedBoldItalic"}
# settings.xml: lo que va antes de <w:embedTrueTypeFonts/> según el esquema
_ANTES_EMBED = ["writeProtection", "view", "zoom", "removePersonalInformation", "removeDateAndTime", "doNotDisplayPageBoundaries",
                "displayBackgroundShape", "printPostScriptOverText", "printFractionalCharacterWidth", "printFormsData"]


def ofuscar_fuente(datos: bytes, clave: str) -> bytes:
    """Ofusca una fuente para incrustarla en un .docx (ECMA-376, 17.8.1): los primeros 32 bytes se combinan con la
    clave (un GUID) leída al revés. La operación es su propia inversa."""
    llave = bytes.fromhex(clave.strip("{}").replace("-", ""))[::-1]
    sal = bytearray(datos)
    for i in range(min(32, len(sal))):
        sal[i] ^= llave[i % 16]
    return bytes(sal)


def _rematar(docx: bytes, usos: Set[Tuple[str, bool, bool]]) -> bytes:
    entrada = zipfile.ZipFile(io.BytesIO(docx))
    archivos = {n: entrada.read(n) for n in entrada.namelist()}
    orden = [n for n in entrada.namelist()]
    # La plantilla trae una miniatura de un documento en blanco: se quita.
    archivos.pop("docProps/thumbnail.jpeg", None)
    if "_rels/.rels" in archivos:
        rels = etree.fromstring(archivos["_rels/.rels"])
        for r in list(rels):
            if r.get("Type") == _REL_MINIATURA:
                rels.remove(r)
        archivos["_rels/.rels"] = etree.tostring(rels, xml_declaration=True, encoding="UTF-8", standalone=True)
    nuevos: List[str] = []
    incrustar = {(f, b, i) for f, b, i in usos if f in _INCRUSTADAS}
    if incrustar:
        tabla = etree.fromstring(archivos["word/fontTable.xml"])
        rels = etree.Element("Relationships", nsmap={None: "http://schemas.openxmlformats.org/package/2006/relationships"})
        por_fuente: Dict[str, list] = {}
        for f, b, i in sorted(incrustar):
            por_fuente.setdefault(f, []).append((b, i))
        n = 0
        for nombre, variantes in por_fuente.items():
            archivos_ttf = FUENTES[_INCRUSTADAS[nombre]][1]
            for viejo in [x for x in tabla if x.get(f"{{{_NS_W}}}name") == nombre]:
                tabla.remove(viejo)
            fuente = etree.SubElement(tabla, f"{{{_NS_W}}}font", {f"{{{_NS_W}}}name": nombre})
            etree.SubElement(fuente, f"{{{_NS_W}}}charset", {f"{{{_NS_W}}}val": "00"})
            etree.SubElement(fuente, f"{{{_NS_W}}}family", {f"{{{_NS_W}}}val": "swiss"})
            etree.SubElement(fuente, f"{{{_NS_W}}}pitch", {f"{{{_NS_W}}}val": "variable"})
            for b, i in sorted(variantes, key=lambda v: _ESTILOS_ARCHIVO[v]):
                n += 1
                clave = "{" + str(uuid.uuid4()).upper() + "}"
                ruta = f"word/fonts/font{n}.odttf"
                archivos[ruta] = ofuscar_fuente((_DIR_FUENTES / archivos_ttf[_ESTILOS_ARCHIVO[(b, i)]]).read_bytes(), clave)
                nuevos.append(ruta)
                etree.SubElement(rels, "Relationship", {"Id": f"rId{n}", "Type": _REL_FUENTE, "Target": f"fonts/font{n}.odttf"})
                etree.SubElement(fuente, f"{{{_NS_W}}}{_ETIQ_INCRUSTAR[(b, i)]}",
                                 {f"{{{_NS_R}}}id": f"rId{n}", f"{{{_NS_W}}}fontKey": clave})
        archivos["word/fontTable.xml"] = etree.tostring(tabla, xml_declaration=True, encoding="UTF-8", standalone=True)
        archivos["word/_rels/fontTable.xml.rels"] = etree.tostring(rels, xml_declaration=True, encoding="UTF-8", standalone=True)
        nuevos.append("word/_rels/fontTable.xml.rels")
        tipos = etree.fromstring(archivos["[Content_Types].xml"])
        ns_t = tipos.nsmap[None]
        if not any(d.get("Extension") == "odttf" for d in tipos):
            tipos.insert(0, etree.Element(f"{{{ns_t}}}Default", Extension="odttf", ContentType=_TIPO_FUENTE))
        archivos["[Content_Types].xml"] = etree.tostring(tipos, xml_declaration=True, encoding="UTF-8", standalone=True)
        ajustes = etree.fromstring(archivos["word/settings.xml"])
        if ajustes.find(f"{{{_NS_W}}}embedTrueTypeFonts") is None:
            lugar = 0
            for k, hijo in enumerate(ajustes):
                if etree.QName(hijo).localname in _ANTES_EMBED:
                    lugar = k + 1
            ajustes.insert(lugar, etree.Element(f"{{{_NS_W}}}embedTrueTypeFonts"))
        archivos["word/settings.xml"] = etree.tostring(ajustes, xml_declaration=True, encoding="UTF-8", standalone=True)
    salida = io.BytesIO()
    with zipfile.ZipFile(salida, "w", zipfile.ZIP_DEFLATED) as z:
        for n in [x for x in orden if x in archivos] + [x for x in nuevos if x not in orden]:
            z.writestr(n, archivos[n])
    return salida.getvalue()


def generar_docx(questions: List[Dict[str, Any]], answer_key: Dict[int, Dict[str, Any]],
                 total_points: float, datos: DatosExamen) -> ResultadoDocx:
    """El Word del examen (y su clave). Mismas preguntas y mismos argumentos que `exportar_pdf.generar_pdf`."""
    plan = construir_plan(questions, answer_key, total_points, datos)
    return ResultadoDocx(docx=_Docx(plan).construir(), preguntas=plan.preguntas, avisos=plan.avisos)
