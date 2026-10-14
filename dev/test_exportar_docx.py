"""
test_exportar_docx.py — Versión 2.8: el examen revisado como Word editable.

    PYTHONUTF8=1 backend/venv/bin/python dev/test_exportar_docx.py

Sin red ni Gemini. No hay Word aquí: se lee el .docx de vuelta (python-docx y XML) y se compara con el plan y con
el PDF del mismo examen:
  1. El marcado del plan (negrita, cursiva, fuente de respaldo, saltos) se lee como ReportLab lo dibuja.
  2. Los 25 exámenes sintéticos y uno con imágenes y fórmulas: cada texto del plan está en el Word, en su orden.
  3. El Word dice lo mismo que el PDF (mismas palabras, mismas letras de opciones, misma clave).
  4. Estructura: secciones, estilos con nombre, numeración «Página X de Y», encabezados, imágenes, tablas.
  5. Fuentes: DejaVu incrustada (y se puede desofuscar a la fuente original); Arial y Times, no.
  6. Los valores raros no rompen nada (nombres con & y <, campos vacíos, tamaños extremos).
"""

import io
import json
import os
import re
import sys
import tempfile
import zipfile
from collections import Counter
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "backend"))
sys.path.insert(0, str(ROOT / "dev"))
sys.path.insert(0, str(ROOT / "dev" / "synthetic"))

TMP = Path(tempfile.mkdtemp(prefix="conv docx "))
os.environ["HOME"] = str(TMP / "home")
os.environ["LOCALAPPDATA"] = str(TMP / "home" / "AppData")
(TMP / "datos" / "data").mkdir(parents=True)
import database  # noqa: E402

database.get_db_path = lambda: TMP / "datos" / "data" / "exams_history.db"  # ANTES de importar main

import docx  # noqa: E402
import pdfplumber  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402
from lxml import etree  # noqa: E402

import xml_fidelity as xf  # noqa: E402  (apaga el logging)
import test_xml_baseline as base  # noqa: E402
import pdf_regresion as pr  # noqa: E402
import exportar_docx as ed  # noqa: E402
import exportar_pdf as ep  # noqa: E402
import main  # noqa: E402
import plan_impreso as pl  # noqa: E402
import seguridad  # noqa: E402
from schema_adapter import adapt  # noqa: E402
from validator import partition_questions  # noqa: E402

fallas = 0


def ok(cond, msg):
    global fallas
    print(("  ok   " if cond else "  FALLA ") + msg)
    fallas += 0 if cond else 1


def norm(s: str) -> str:
    return " ".join(str(s).replace("\u00a0", " ").replace("\t", " ").split())


def datos(**kw):
    return ep.DatosExamen(**{**pr.DATOS, **kw})


def abrir(blob: bytes) -> "docx.document.Document":
    return docx.Document(io.BytesIO(blob))


# ══════════════════════════════════════════════════════════════════════════
print("1. El marcado del plan, leído como lo dibuja ReportLab")
# ══════════════════════════════════════════════════════════════════════════
t = ed.leer_marcado("<b>A)</b> Texto   con   espacios  y <i>cursiva</i><br/>otra línea &amp; &lt;etiqueta&gt;")
ok([x.texto for x in t if x.texto is not None] == ["A)", " Texto con espacios y ", "cursiva", "otra línea & <etiqueta>"],
   f"los blancos se juntan en uno y las entidades se leen ({[x.texto for x in t]})")
ok(t[0].negrita and not t[1].negrita and t[2].cursiva and any(x.texto is None for x in t), "negrita, cursiva y salto de línea")
t = ed.leer_marcado("  hola\u00a0\u00a0mundo  ")
ok([x.texto for x in t] == ["hola\u00a0\u00a0mundo"], f"el espacio de no separación NO se junta y los extremos se recortan ({[x.texto for x in t]!r})")
t = ed.leer_marcado('Si <font name="DejaVu">∀</font>x <font size=\'8.0\' color=\'#555555\'>(2 pts)</font>')
ok(t[1].fuente == "DejaVu" and t[3].tam == 8.0 and t[3].color == "#555555" and t[2].fuente is None, "fuente de respaldo, tamaño y color")
ok(ed.leer_marcado("") == [] and ed.leer_marcado("   ") == [], "marcado vacío")
llave = "{0F1E2D3C-4B5A-6978-8796-A5B4C3D2E1F0}"
datos_ttf = (pl._DIR_FUENTES / "DejaVuSans.ttf").read_bytes()
ofus = ed.ofuscar_fuente(datos_ttf, llave)
ok(ofus[:32] != datos_ttf[:32] and ofus[32:] == datos_ttf[32:] and ed.ofuscar_fuente(ofus, llave) == datos_ttf,
   "la ofuscación de la fuente cambia solo los 32 primeros bytes y se deshace")


# ══════════════════════════════════════════════════════════════════════════
def textos_plan(bloques, salida=None):
    """Los textos de los párrafos que el plan pide (sin las hojas de respuestas, que van con tabuladores)."""
    salida = [] if salida is None else salida
    for b in bloques:
        if isinstance(b, pl.Parrafo):
            salida.append(norm((b.vineta + " " if b.vineta else "") + "".join(" " if x.texto is None else x.texto for x in ed.leer_marcado(b.marcado))))
        elif isinstance(b, pl.Opcion):
            salida.append(norm("".join(x.texto or " " for x in ed.leer_marcado(b.marcado))))
        elif isinstance(b, (pl.Sangrado, pl.Grupo)):
            textos_plan(b.hijos, salida)
        elif isinstance(b, pl.TablaPareo):
            for izq, der in b.filas:
                textos_plan([izq, der], salida)
        elif isinstance(b, pl.TablaHuecos):
            for num, celda in b.filas:
                textos_plan([num] + celda, salida)
        elif isinstance(b, pl.LogosCabecera):
            textos_plan(b.centro, salida)
        elif isinstance(b, pl.CuadroEstudiante):
            textos_plan([c for fila in b.filas for c in fila if isinstance(c, pl.Parrafo)], salida)
        elif isinstance(b, pl.FilaDocente):
            textos_plan([x for x in (b.izquierda, b.nota) if x is not None], salida)
    return [s for s in salida if s]


def parrafos_docx(d):
    """Todos los párrafos del Word (cuerpo y celdas) en orden, con su estilo."""
    salida = []
    for hijo in d.element.body.iterchildren():
        if hijo.tag == qn("w:p"):
            salida.append(docx.text.paragraph.Paragraph(hijo, d))
        elif hijo.tag == qn("w:tbl"):
            for p in hijo.iter(qn("w:p")):
                salida.append(docx.text.paragraph.Paragraph(p, d))
    return salida


ESTILOS_APARTE = {"Cátedra respuestas", "Cátedra burbujas"}


# ══════════════════════════════════════════════════════════════════════════
print("2. Los 25 exámenes sintéticos y uno con imágenes y fórmulas: lo que pide el plan está en el Word")
# ══════════════════════════════════════════════════════════════════════════
bancos = {}
for nombre, qs in base.bancos().items():
    questions, key = adapt(xf.perfect_model(qs))
    validas, _ = partition_questions(questions, key)
    if validas:
        bancos[nombre] = (validas, key)
bancos["extra"] = pr.examen_extra()
total_docx = 0
for nombre, (qs, key) in bancos.items():
    problemas = []
    for variante in ("examen_y_clave", "folleto"):
        d_ex = datos(**pr.VARIANTES[variante])
        plan = pl.construir_plan(qs, key, 20, d_ex)
        r = ed.generar_docx(qs, key, 20, d_ex)
        total_docx += 1
        d = abrir(r.docx)
        esperado = Counter(s for sec in plan.secciones for s in textos_plan(sec.bloques))
        real = Counter(norm(p.text) for p in parrafos_docx(d) if norm(p.text) and p.style.name not in ESTILOS_APARTE)
        # lo que no está o sobra
        falta = esperado - real
        sobra = real - esperado
        if falta:
            problemas.append(f"{variante}: faltan {list(falta.items())[:2]}")
        if sobra:
            problemas.append(f"{variante}: sobran {list(sobra.items())[:2]}")
        if r.preguntas != plan.preguntas or r.avisos != plan.avisos:
            problemas.append(f"{variante}: preguntas/avisos distintos")
    ok(not problemas, f"{nombre}: texto del Word = texto del plan {problemas or ''}")
ok(total_docx >= 50, f"se generaron {total_docx} Word")


# ══════════════════════════════════════════════════════════════════════════
print("3. El Word dice lo mismo que el PDF")
# ══════════════════════════════════════════════════════════════════════════
_EMOJI = re.compile("[\U00010000-\U0010FFFF\u1f60]")      # pdfplumber lee mal los emojis: no se comparan
_NBSP = chr(0xA0)


def _contar(texto: str, c: Counter) -> None:
    for tok in _EMOJI.sub("", texto.replace(_NBSP, " ").replace("\t", " ")).split():
        c[tok] += 1


def palabras_pdf(pdf: bytes, plan) -> Counter:
    """Las palabras del PDF sin «Página X de Y» ni el encabezado corrido (el Word los lleva en su encabezado y pie)."""
    corridos = {norm(f"{s.izquierda} {s.derecha}") for s in plan.secciones}
    c = Counter()
    with pdfplumber.open(io.BytesIO(pdf)) as doc:
        for pag in doc.pages:
            for linea in (pag.extract_text() or "").splitlines():
                if re.fullmatch(r"Página \d+ de \d+", linea.strip()) or norm(linea) in corridos:
                    continue
                _contar(linea, c)
    return c


def letras_de_burbujas(bloques, c=None) -> Counter:
    """Las letras que el PDF escribe DENTRO de las burbujas de la hoja de respuestas (en el Word son imagen)."""
    c = Counter() if c is None else c
    for b in bloques:
        if isinstance(b, pl.RejillaBurbujas):
            for fila in b.celdas:
                for bur in fila:
                    if bur is not None:
                        c.update(bur.letras)
        elif isinstance(b, (pl.Grupo, pl.Sangrado)):
            letras_de_burbujas(b.hijos, c)
    return c


def palabras_docx(blob: bytes) -> Counter:
    d = abrir(blob)
    c = Counter()
    for p in parrafos_docx(d):
        _contar(p.text, c)
    return c



for nombre in ("extra",) + tuple(list(bancos)[:6]):
    qs, key = bancos[nombre]
    for variante in ("examen_y_clave", "folleto"):
        d_ex = datos(**pr.VARIANTES[variante])
        plan_x = pl.construir_plan(qs, key, 20, d_ex)
        pdf_w = palabras_pdf(ep.generar_pdf(qs, key, 20, d_ex).pdf, plan_x)
        for sec in plan_x.secciones:
            pdf_w = pdf_w - letras_de_burbujas(sec.bloques)
        doc_w = palabras_docx(ed.generar_docx(qs, key, 20, d_ex).docx)
        falta, sobra = pdf_w - doc_w, doc_w - pdf_w
        ok(not falta and not sobra, f"{nombre} / {variante}: mismas palabras que el PDF "
           f"{'' if not (falta or sobra) else (list(falta.items())[:4], list(sobra.items())[:4])}")


# ══════════════════════════════════════════════════════════════════════════
print("4. Estructura del documento")
# ══════════════════════════════════════════════════════════════════════════
qs, key = bancos["extra"]
r = ed.generar_docx(qs, key, 20, datos(contenido="folleto_hoja_clave"))
d = abrir(r.docx)
ok(len(d.sections) == 3, f"folleto + hoja + clave = 3 secciones ({len(d.sections)})")
r1 = ed.generar_docx(qs, key, 20, datos(contenido="solo_examen"))
ok(len(abrir(r1.docx).sections) == 1, "solo el examen = 1 sección")
z = zipfile.ZipFile(io.BytesIO(r.docx))
for nombre_xml in z.namelist():
    if nombre_xml.endswith((".xml", ".rels")):
        etree.fromstring(z.read(nombre_xml))
ok(True, f"todas las partes XML son válidas ({len(z.namelist())} archivos)")
ok("docProps/thumbnail.jpeg" not in z.namelist(), "se quitó la miniatura de la plantilla")
estilos = {s.name for s in d.styles if s.name.startswith("Cátedra")}
ok({"Cátedra enunciado", "Cátedra parte", "Cátedra rotulo", "Cátedra cabecera", "Cátedra burbujas"} <= estilos, f"estilos con nombre para el docente: {len(estilos)}")
xml_doc = z.read("word/document.xml").decode()
for sec_i, sec in enumerate(d.sections):
    ok(sec._sectPr.find(qn("w:pgNumType")) is not None and sec._sectPr.find(qn("w:pgNumType")).get(qn("w:start")) == "1",
       f"sección {sec_i + 1}: la numeración empieza en 1")
    ok(sec.footer.paragraphs[0]._p.xpath('.//w:instrText[contains(., "SECTIONPAGES")]') != [], f"sección {sec_i + 1}: «Página X de Y» con campos de Word")
ok(d.sections[0].different_first_page_header_footer and not d.sections[2].different_first_page_header_footer, "el encabezado corrido salta la primera hoja del examen, no la de la clave")
ok("CLAVE DE RESPUESTAS" in d.sections[2].header.paragraphs[0].text, f"la clave lleva su encabezado ({d.sections[2].header.paragraphs[0].text!r})")
ok(len(d.inline_shapes) >= 20, f"casillas y burbujas son imágenes ({len(d.inline_shapes)})")
g = pl.construir_plan(qs, key, 20, datos()).geo
ok(abs(d.sections[0].page_width.pt - g.papel[0]) < 0.5 and abs(d.sections[0].left_margin.pt - g.m_izq) < 0.5, "papel y márgenes como en el PDF")
ok(d.core_properties.title and d.core_properties.last_modified_by == "Cátedra", "propiedades del documento")
ok(qn("w:embedTrueTypeFonts") in {c.tag for c in etree.fromstring(z.read("word/settings.xml"))}, "settings.xml pide las fuentes incrustadas")


# ══════════════════════════════════════════════════════════════════════════
print("5. Fuentes: DejaVu incrustada; Arial y Times con su nombre")
# ══════════════════════════════════════════════════════════════════════════
def incrustadas(blob: bytes):
    z = zipfile.ZipFile(io.BytesIO(blob))
    return [n for n in z.namelist() if n.startswith("word/fonts/")], z


nombres, z = incrustadas(ed.generar_docx(qs, key, 20, datos()).docx)
ok(len(nombres) in (2, 3, 4), f"con DejaVu se incrustan las variantes que se usan ({len(nombres)})")
tabla = etree.fromstring(z.read("word/fontTable.xml"))
fuente = [f for f in tabla if f.get(qn("w:name")) == "DejaVu Sans"][0]
claves = [(e.get(qn("w:fontKey")), e.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")) for e in fuente if "embed" in etree.QName(e).localname]
rels = etree.fromstring(z.read("word/_rels/fontTable.xml.rels"))
destinos = {r.get("Id"): r.get("Target") for r in rels}
bien = 0
originales = {open_ttf: (pl._DIR_FUENTES / open_ttf).read_bytes() for open_ttf in pl.FUENTES["dejavu"][1]}
for clave, rid in claves:
    des = ed.ofuscar_fuente(z.read("word/" + destinos[rid]), clave)
    bien += any(des == o for o in originales.values())
ok(bien == len(claves) and len(claves) == len(nombres), f"cada fuente incrustada se desofusca a la DejaVu original ({bien}/{len(claves)})")
ok(any(d.get("Extension") == "odttf" for d in etree.fromstring(z.read("[Content_Types].xml"))), "tipo de contenido de las fuentes ofuscadas")

arial = ed.generar_docx(qs, key, 20, datos(fuente_titulos="arial", fuente_preguntas="arial"))
nom_a, z_a = incrustadas(arial.docx)
xml_a = z_a.read("word/styles.xml").decode()
ok("Arial" in xml_a and "Liberation" not in xml_a, "Arial se escribe como «Arial» (no como la copia del PDF)")
ok(len(nom_a) >= 1 and len(nom_a) < len(nombres) + 1, f"con Arial solo se incrusta DejaVu para lo que Arial no tiene ({len(nom_a)})")
times = ed.generar_docx(qs, key, 20, datos(fuente_titulos="times", fuente_preguntas="times"))
ok("Times New Roman" in z_a.read("word/styles.xml").decode() or "Times New Roman" in incrustadas(times.docx)[1].read("word/styles.xml").decode(), "Times New Roman")
sin_nada, _ = incrustadas(ed.generar_docx([bancos["extra"][0][1]], {2: bancos["extra"][1][2]}, 20, datos(fuente_titulos="arial", fuente_preguntas="arial")).docx)
ok(sin_nada == [], "un examen que Arial escribe entero no incrusta nada")


# ══════════════════════════════════════════════════════════════════════════
print("6. Valores raros")
# ══════════════════════════════════════════════════════════════════════════
raros = [
    dict(institucion="A & B <C>", materia='Comillas "dobles" y \'simples\'', docente="Ñandú Ünico", grupo="", fecha=""),
    dict(institucion="", facultad="", departamento="", materia="", actividad="", docente="", instrucciones=""),
    dict(tam_titulos=20, tam_preguntas=7, renglones_ensayo=24),
    dict(tam_titulos=7, tam_preguntas=20, papel="legal", margenes="anchos"),
    dict(institucion="X" * 160, docente="Y" * 160, campos_estudiante=False),
    dict(logo_izquierdo="no es una imagen", logo_derecho="aGVsbG8="),
    dict(puntos_por_pregunta=False, partes=False, mezclar=False, contenido="solo_clave"),
    dict(version="B", rotulo_docente="profesor", contenido="folleto_hoja_clave", fuente_titulos="times", fuente_preguntas="arial"),
]
for i, kw in enumerate(raros, 1):
    try:
        res = ed.generar_docx(qs, key, 20, datos(**kw))
        abrir(res.docx)
        ok(True, f"caso {i}: {sorted(kw)[:4]}… sale y se vuelve a abrir ({len(res.docx) // 1024} KB, avisos {len(res.avisos)})")
    except Exception as exc:  # noqa: BLE001
        ok(False, f"caso {i}: {kw} → {type(exc).__name__}: {exc}")
for malo in (dict(contenido="otro"), dict(fuente_titulos="comic")):
    try:
        ed.generar_docx(qs, key, 20, datos(**malo))
        ok(False, f"{malo} debería rechazarse")
    except ValueError:
        ok(True, f"{malo} se rechaza con ValueError, como el PDF")
try:
    ed.generar_docx(qs, key, 20, datos(puntos_enteros=True))
    ok(False, "puntos con decimales + «puntos enteros» debería rechazarse")
except ep.PuntosNoEnteros:
    ok(True, "«puntos enteros» con decimales se rechaza como en el PDF (nunca se redondea)")

# ══════════════════════════════════════════════════════════════════════════
print("7. Las rutas: /api/exportar_examen y las versiones en Word")
# ══════════════════════════════════════════════════════════════════════════
main.init_db()
cli = TestClient(main.app, base_url="http://127.0.0.1:8000")
H = {"X-Conversor-Token": seguridad.TOKEN}
qs8, key8 = bancos["extra"]
cuerpo = {"filename": "mi examen.docx", "total_points": 20, "questions": qs8, "answer_key": {str(k): v for k, v in key8.items()},
          "datos": {"docente": "Ana", "materia": "Física", "formato": "word"}}
r = cli.post("/api/exportar_examen", json=cuerpo, headers=H)
ok(r.status_code == 200, f"Word: 200 ({r.status_code})")
ok(r.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.wordprocessingml.document") and r.content[:2] == b"PK",
   f"…con un .docx ({r.headers['content-type']})")
ok(".docx" in r.headers["content-disposition"] and "examen_y_clave" in r.headers["content-disposition"], f"nombre {r.headers['content-disposition']}")
ok(len(abrir(r.content).sections) == 2, "el .docx descargado se abre y trae examen y clave")
info = json.loads(r.headers["x-pdf-info"])
ok(info["preguntas"] == 8 and info["paginas"] is None, f"cabecera X-PDF-Info del Word {info}")
r_pdf = cli.post("/api/exportar_examen", json={**cuerpo, "datos": {**cuerpo["datos"], "formato": "pdf"}}, headers=H)
ok(r_pdf.status_code == 200 and r_pdf.content[:5] == b"%PDF-" and ".pdf" in r_pdf.headers["content-disposition"], "formato pdf: el PDF de siempre")
ok(json.loads(r_pdf.headers["x-pdf-info"])["paginas"] >= 4, "…y dice cuántas páginas tiene")
r_amb = cli.post("/api/exportar_examen", json={**cuerpo, "datos": {**cuerpo["datos"], "formato": "ambos"}}, headers=H)
zamb = zipfile.ZipFile(io.BytesIO(r_amb.content))
ok(r_amb.status_code == 200 and r_amb.headers["content-type"] == "application/zip" and "pdf_y_word" in r_amb.headers["content-disposition"],
   f"«ambos»: un ZIP ({r_amb.headers['content-disposition']})")
ok(sorted(n.rsplit(".", 1)[1] for n in zamb.namelist()) == ["docx", "pdf"] and zamb.namelist()[0].rsplit(".", 1)[0] == zamb.namelist()[1].rsplit(".", 1)[0],
   f"…con el PDF y el Word del mismo nombre {zamb.namelist()}")
ok(zamb.read([n for n in zamb.namelist() if n.endswith(".pdf")][0])[:5] == b"%PDF-" and abrir(zamb.read([n for n in zamb.namelist() if n.endswith(".docx")][0])), "…y los dos se abren")
ok(cli.post("/api/exportar_examen", json=cuerpo).status_code == 401, "sin token: 401")
ok(cli.post("/api/exportar_examen", json={**cuerpo, "datos": {"formato": "excel"}}, headers=H).status_code == 422, "formato desconocido: 422")
ok(cli.post("/api/exportar_examen", json={**cuerpo, "datos": {"formato": "word", "fuente_titulos": "comic"}}, headers=H).status_code == 422, "tipo de letra desconocido: 422")
ok(cli.post("/api/exportar_examen", json={**cuerpo, "datos": {"formato": "word", "contenido": "otra cosa"}}, headers=H).status_code == 422, "contenido desconocido: 422")
malo = {**cuerpo, "questions": [{"num": 1, "type": "inventada", "data": {}}], "answer_key": {"1": {"type": "inventada", "answer": "A"}}}
ok(cli.post("/api/exportar_examen", json=malo, headers=H).status_code == 422, "la misma validación que el XML: tipo desconocido -> 422")
r_dec = cli.post("/api/exportar_examen", json={**cuerpo, "total_points": 20.5, "datos": {"formato": "word", "puntos_enteros": True}}, headers=H)
ok(r_dec.status_code == 422 and "decimales" in r_dec.json()["detail"], "«puntos enteros» con decimales: 422 también para Word (nunca se redondea)")
ok(cli.post("/api/exportar_pdf", json={**cuerpo, "datos": {"formato": "word"}}, headers=H).content[:5] == b"%PDF-", "/api/exportar_pdf sigue dando PDF")

# versiones: un archivo por versión y por formato
vcuerpo = {"filename": "banco.docx", "total_points": 20, "questions": qs8, "answer_key": cuerpo["answer_key"], "datos": {"formato": "word", "contenido": "solo_examen"},
           "versiones": [{"etiqueta": "A", "nums": [1, 2, 3, 4], "total_points": 10}, {"etiqueta": "B", "nums": [5, 6, 7, 8], "total_points": 10}]}
rv = cli.post("/api/exportar_versiones", json=vcuerpo, headers=H)
zv = zipfile.ZipFile(io.BytesIO(rv.content)) if rv.status_code == 200 else None
ok(rv.status_code == 200 and zv is not None and sorted(zv.namelist()) == ["banco_version_A_examen.docx", "banco_version_B_examen.docx"],
   f"versiones en Word: un .docx por versión ({zv.namelist() if zv else rv.text[:100]})")
va = abrir(zv.read("banco_version_A_examen.docx"))
texto_va = " ".join(p.text for p in parrafos_docx(va))
ok("VERSIÓN A" in texto_va and "Desarrolla el tema" not in texto_va and "Mira la imagen" in texto_va, "cada Word trae SOLO sus preguntas y su nombre de versión")
rv2 = cli.post("/api/exportar_versiones", json={**vcuerpo, "datos": {**vcuerpo["datos"], "formato": "ambos"}}, headers=H)
zv2 = zipfile.ZipFile(io.BytesIO(rv2.content))
ok(sorted(zv2.namelist()) == ["banco_version_A_examen.docx", "banco_version_A_examen.pdf", "banco_version_B_examen.docx", "banco_version_B_examen.pdf"],
   f"«ambos»: un PDF y un Word por versión ({zv2.namelist()})")
ok(len(json.loads(rv2.headers["x-pdf-info"])["versiones"]) == 2, "…y el resumen de versiones")
rv3 = cli.post("/api/exportar_versiones", json={**vcuerpo, "total_points": 20, "datos": {"formato": "word", "puntos_enteros": True}, "versiones": [{"etiqueta": "A", "nums": [1, 2, 3], "total_points": 7}]}, headers=H)
ok(rv3.status_code == 422 and rv3.json()["detail"].startswith("Versión A:"), f"un decimal en una versión se rechaza con su nombre ({rv3.status_code})")

print()
print("TODO OK" if fallas == 0 else f"{fallas} FALLA(S)")
sys.exit(1 if fallas else 0)
